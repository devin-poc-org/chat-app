"""
File Handler Module
Extracts text from various file types (txt, code files, PDF, Word docs)
"""

import os
import chardet
from typing import Optional, Tuple
from PyPDF2 import PdfReader
from docx import Document


class FileHandler:
    """Handle file uploads and text extraction"""
    
    MAX_FILE_SIZE = 1 * 1024 * 1024  # 1MB
    
    SUPPORTED_TEXT_EXTENSIONS = {
        '.txt', '.md', '.py', '.js', '.java', '.cpp', '.c', '.h', '.hpp',
        '.cs', '.rb', '.go', '.rs', '.php', '.html', '.css', '.xml', '.json',
        '.yaml', '.yml', '.sh', '.bash', '.sql', '.r', '.swift', '.kt', '.ts',
        '.jsx', '.tsx', '.vue', '.scala', '.pl', '.lua', '.m', '.mm', '.gradle',
        '.properties', '.conf', '.config', '.ini', '.toml', '.log'
    }
    
    @staticmethod
    def validate_file(file_data: bytes, filename: str) -> Tuple[bool, Optional[str]]:
        """Validate file size and type"""
        if len(file_data) > FileHandler.MAX_FILE_SIZE:
            return False, f"File size exceeds 1MB limit (size: {len(file_data) / 1024 / 1024:.2f}MB)"
        
        ext = os.path.splitext(filename)[1].lower()
        if ext not in FileHandler.SUPPORTED_TEXT_EXTENSIONS and ext not in ['.pdf', '.docx', '.doc']:
            return False, f"Unsupported file type: {ext}"
        
        return True, None
    
    @staticmethod
    def detect_encoding(file_data: bytes) -> str:
        """Detect file encoding"""
        result = chardet.detect(file_data)
        return result['encoding'] or 'utf-8'
    
    @staticmethod
    def extract_text_from_pdf(file_data: bytes) -> str:
        """Extract text from PDF file"""
        try:
            import io
            pdf_file = io.BytesIO(file_data)
            reader = PdfReader(pdf_file)
            
            text_parts = []
            for page_num, page in enumerate(reader.pages, 1):
                text = page.extract_text()
                if text.strip():
                    text_parts.append(f"--- Page {page_num} ---\n{text}")
            
            return "\n\n".join(text_parts)
        except Exception as e:
            raise Exception(f"Failed to extract text from PDF: {str(e)}")
    
    @staticmethod
    def extract_text_from_docx(file_data: bytes) -> str:
        """Extract text from Word document"""
        try:
            import io
            docx_file = io.BytesIO(file_data)
            doc = Document(docx_file)
            
            text_parts = []
            for para in doc.paragraphs:
                if para.text.strip():
                    text_parts.append(para.text)
            
            for table in doc.tables:
                for row in table.rows:
                    row_text = ' | '.join(cell.text.strip() for cell in row.cells)
                    if row_text.strip():
                        text_parts.append(row_text)
            
            return "\n\n".join(text_parts)
        except Exception as e:
            raise Exception(f"Failed to extract text from Word document: {str(e)}")
    
    @staticmethod
    def extract_text_from_file(file_data: bytes, filename: str) -> Tuple[str, str]:
        """
        Extract text from file based on file type
        Returns: (extracted_text, file_type)
        """
        ext = os.path.splitext(filename)[1].lower()
        
        try:
            if ext == '.pdf':
                text = FileHandler.extract_text_from_pdf(file_data)
                return text, 'PDF'
            
            elif ext in ['.docx', '.doc']:
                text = FileHandler.extract_text_from_docx(file_data)
                return text, 'Word Document'
            
            elif ext in FileHandler.SUPPORTED_TEXT_EXTENSIONS:
                encoding = FileHandler.detect_encoding(file_data)
                text = file_data.decode(encoding, errors='replace')
                return text, 'Text File'
            
            else:
                raise Exception(f"Unsupported file type: {ext}")
        
        except Exception as e:
            raise Exception(f"Failed to extract text from {filename}: {str(e)}")
    
    @staticmethod
    def format_file_content_for_chat(filename: str, file_type: str, content: str, max_preview_length: int = 500) -> str:
        """Format file content for display in chat"""
        preview = content[:max_preview_length]
        if len(content) > max_preview_length:
            preview += "..."
        
        return f"📎 **File uploaded: {filename}** ({file_type})\n\n```\n{preview}\n```"
    
    @staticmethod
    def format_file_content_for_model(filename: str, content: str) -> str:
        """Format file content for sending to model"""
        return f"[File: {filename}]\n\n{content}\n\n[End of file]"
